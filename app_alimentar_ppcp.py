"""
Alimentador PPCP - lê PDFs de Relatório de Atingimento por Setor/Peças
 e alimenta a aba 5_ACOMPANHAMENTO da planilha mãe.

Como usar com interface:
    python app_alimentar_ppcp.py

Como usar sem interface:
    python app_alimentar_ppcp.py --sem-gui --planilha PPCP_ACOMPAN.xlsx --pdfs "SEC 1 - 10180,10148.pdf" --saida PPCP_ACOMPAN_atualizada.xlsx --data 30/06/2026

Dependências:
    pip install pymupdf openpyxl
"""

from __future__ import annotations

import argparse
import copy
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple

try:
    import fitz  # PyMuPDF
except ImportError as exc:
    raise SystemExit(
        "Falta instalar o PyMuPDF. Rode: pip install pymupdf"
    ) from exc

try:
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
except ImportError as exc:
    raise SystemExit(
        "Falta instalar o openpyxl. Rode: pip install openpyxl"
    ) from exc


ABA_PADRAO = "5_ACOMPANHAMENTO"

# Mapeamento das colunas na aba 5_ACOMPANHAMENTO.
COL_DATA_PROGRAMADA = 1     # A
COL_OP_LOTE = 2             # B
COL_CLIENTE_PEDIDO = 3      # C - opcional: recebe o produto do cabeçalho do PDF, ex.: MAD 0707 (VERMELHO)
COL_PROD_EQUIP = 4          # D
COL_CODIGO_PECA = 5         # E
COL_DESCRICAO_PECA = 6      # F
COL_OPERACAO = 7            # G
COL_SETOR = 8               # H
COL_MAQUINA_POSTO = 9       # I
COL_QTDE_PROGRAMADA = 10    # J

# Colunas com fórmulas que devem ser copiadas/traduzidas para a nova linha.
# O programa procura uma linha-modelo com fórmulas e traduz automaticamente a referência da linha.
COLUNAS_FORMULAS_PREFERIDAS = [11, 12, 13, 18, 20, 22, 23, 24, 25, 26, 27, 28, 29, 30]


@dataclass
class RegistroPDF:
    arquivo: str
    pagina: int
    op_lote: str
    cliente_pedido: str
    produto_equipamento: str
    codigo_peca: str
    descricao_peca: str
    operacao: str
    setor: str
    maquina_posto: str
    qtde_programada: float


def limpar_texto(valor: object) -> str:
    return re.sub(r"\s+", " ", str(valor or "").strip())


def numero_br_para_float(valor: str) -> float:
    """Converte 2.590,00000 ou 820,00000 para float."""
    valor = limpar_texto(valor)
    valor = valor.replace(".", "").replace(",", ".")
    return float(valor)


def float_para_int_se_possivel(valor: float) -> float | int:
    return int(valor) if abs(valor - int(valor)) < 0.000001 else valor


def parse_data_opcional(valor: str) -> Optional[datetime]:
    valor = limpar_texto(valor)
    if not valor:
        return None
    formatos = ["%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"]
    for fmt in formatos:
        try:
            return datetime.strptime(valor, fmt)
        except ValueError:
            pass
    raise ValueError("Data Programada inválida. Use DD/MM/AAAA, DD/MM/AA ou AAAA-MM-DD.")


def linhas_texto_pagina(page: fitz.Page) -> List[str]:
    texto = page.get_text("text") or ""
    return [limpar_texto(l) for l in texto.splitlines() if limpar_texto(l)]


def proxima_linha_nao_vazia(linhas: Sequence[str], indice: int) -> str:
    for j in range(indice + 1, len(linhas)):
        if linhas[j]:
            return linhas[j]
    return ""


def lote_valido(valor: str) -> bool:
    """Lote válido deve ter pelo menos 4 dígitos; evita capturar número da página."""
    valor = limpar_texto(valor)
    return bool(re.search(r"\d{4,}", valor))


def setor_valido(valor: str) -> bool:
    valor = limpar_texto(valor).upper()
    if not valor:
        return False
    termos_invalidos = ["TOTAL", "PÁGINA", "PAGINA", "EMISSÃO", "EMISSAO", "RELATÓRIO", "RELATORIO"]
    return not any(t in valor for t in termos_invalidos)


def extrair_lotes_nome_arquivo(caminho_pdf: str | Path) -> str:
    """Fallback: tenta pegar lotes do nome do arquivo, ex.: SEC 1 - 10180,10148.pdf."""
    nome = Path(caminho_pdf).stem
    candidatos = re.findall(r"\b\d{4,6}(?:\s*[,;]\s*\d{4,6})*\b", nome)
    return limpar_texto(candidatos[-1].replace(";", ",")) if candidatos else ""


def agrupar_palavras_por_linha_visual(page: fitz.Page, tolerancia: float = 4.0) -> List[List[tuple]]:
    words = sorted(page.get_text("words") or [], key=lambda w: (w[1], w[0]))
    linhas: List[List[tuple]] = []
    for w in words:
        cy = (w[1] + w[3]) / 2
        colocado = False
        for linha in linhas:
            cy_linha = sum((x[1] + x[3]) / 2 for x in linha) / len(linha)
            if abs(cy - cy_linha) <= tolerancia:
                linha.append(w)
                colocado = True
                break
        if not colocado:
            linhas.append([w])
    return [sorted(l, key=lambda w: w[0]) for l in linhas]


def extrair_cabecalho_por_posicao(page: fitz.Page) -> Tuple[str, str, str]:
    """Extrai cabeçalho por coordenadas, evitando confundir Página/Total com lote/setor."""
    lotes = ""
    produto = ""
    setor = ""

    for linha_words in agrupar_palavras_por_linha_visual(page):
        textos = [w[4] for w in linha_words]
        texto_linha = limpar_texto(" ".join(textos))
        texto_upper = texto_linha.upper()

        # Linha visual dos lotes. No PDF, o rótulo fica à esquerda, o valor logo à direita,
        # e o produto fica bem mais à direita.
        if any(t.upper().startswith("LOTES:") for t in textos):
            rotulos = [w for w in linha_words if w[4].upper().startswith("LOTES:")]
            x_rotulo = rotulos[0][2] if rotulos else 0
            lote_words = []
            produto_words = []
            for w in linha_words:
                palavra = w[4]
                if palavra.upper().startswith("LOTES:"):
                    continue
                if w[0] > x_rotulo and w[0] < 230 and re.fullmatch(r"[0-9,.;/\-]+", palavra):
                    lote_words.append(palavra)
                elif w[0] > 230 and not any(palavra.upper().startswith(x) for x in ["EQUIPAMENTO:", "FAMILIA:", "PÁGINA:", "PAGINA:"]):
                    produto_words.append(palavra)
            if lote_words:
                lotes = limpar_texto("".join(lote_words).replace(";", ","))
            if produto_words:
                produto = limpar_texto(" ".join(produto_words))

        # Linha visual do setor. Importante: não aceitar "Total do Setor".
        # O rótulo correto fica no cabeçalho, à esquerda e antes da tabela.
        if ("SETOR:" in texto_upper) and ("TOTAL" not in texto_upper):
            rotulos = [w for w in linha_words if w[4].upper() == "SETOR:"]
            if rotulos and rotulos[0][0] < 120 and rotulos[0][1] < 160:
                x_rotulo = rotulos[0][2]
                setor_words = [w[4] for w in linha_words if w[0] > x_rotulo and w[4].upper() != "SETOR:"]
                setor = limpar_texto(" ".join(setor_words))

    return lotes, produto, setor


def extrair_cabecalho(page: fitz.Page) -> Tuple[str, str, str]:
    """Retorna (lotes, produto_cabecalho, setor/equipamento).

    A primeira versão lia pela ordem textual do PyMuPDF. Em alguns PDFs essa ordem vem assim:
    Página -> Lotes -> Produto, fazendo o número da página cair na coluna lote.
    Agora a leitura prioriza a posição visual do cabeçalho e valida o resultado.
    """
    lotes, produto_cabecalho, setor = extrair_cabecalho_por_posicao(page)
    linhas = linhas_texto_pagina(page)

    # Fallback textual para PDFs em que a posição não venha bem estruturada.
    for i, linha in enumerate(linhas):
        up = linha.upper()
        if up.startswith("LOTES:") and not lote_valido(lotes):
            resto = limpar_texto(linha.split(":", 1)[1]) if ":" in linha else ""
            if resto:
                m = re.match(r"([0-9,.;/\-\s]+)(.*)$", resto)
                if m:
                    candidato_lote = limpar_texto(m.group(1).replace(";", ","))
                    if lote_valido(candidato_lote):
                        lotes = candidato_lote
                    if not produto_cabecalho:
                        produto_cabecalho = limpar_texto(m.group(2))
            else:
                # Procura primeiro uma linha próxima que pareça lote real.
                janela = linhas[max(0, i - 3): min(len(linhas), i + 4)]
                for cand in janela:
                    if lote_valido(cand) and re.fullmatch(r"[0-9,.;/\-\s]+", cand):
                        lotes = limpar_texto(cand.replace(";", ","))
                        break
                # Produto costuma estar depois do rótulo dos lotes.
                for cand in linhas[i + 1: min(len(linhas), i + 5)]:
                    if cand and not cand.upper().startswith(("EQUIPAMENTO", "FAMILIA", "SETOR", "PÁGINA", "PAGINA")) and not lote_valido(cand):
                        produto_cabecalho = produto_cabecalho or cand
                        break

        # Não usar "Total do Setor" como setor.
        if up.startswith("SETOR:") and "TOTAL" not in up and not setor_valido(setor):
            resto = limpar_texto(linha.split(":", 1)[1]) if ":" in linha else ""
            if setor_valido(resto):
                setor = resto
            else:
                # Exemplo de extração textual: Setor: / SECCIONADORA 1 / 46.
                proximas = [x for x in linhas[i + 1: min(len(linhas), i + 5)] if x]
                nums = [x for x in proximas if re.fullmatch(r"\d{1,3}", x)]
                textos = [x for x in proximas if not re.fullmatch(r"\d{1,3}", x) and setor_valido(x)]
                if nums and textos:
                    setor = f"{nums[0]} {textos[0]}"
                elif textos:
                    setor = textos[0]

    lotes = limpar_texto(lotes).replace(";", ",")
    produto_cabecalho = limpar_texto(produto_cabecalho)
    setor = limpar_texto(setor)

    # Último fallback por palavras, agora protegido contra "Total do Setor".
    if not setor_valido(setor):
        setor = extrair_setor_por_palavras(page)

    if not setor_valido(setor):
        setor = ""

    return lotes, produto_cabecalho, setor


def extrair_setor_por_palavras(page: fitz.Page) -> str:
    for ws in agrupar_palavras_por_linha_visual(page):
        texto = limpar_texto(" ".join(w[4] for w in ws))
        texto_upper = texto.upper()
        if "SETOR:" in texto_upper and "TOTAL" not in texto_upper:
            tokens = [w for w in ws if w[4].upper() == "SETOR:"]
            if not tokens:
                continue
            rotulo = tokens[0]
            # Cabeçalho fica no topo/esquerda. Evita a linha "Total do Setor" no rodapé da tabela.
            if rotulo[0] > 120 or rotulo[1] > 160:
                continue
            setor_words = [w[4] for w in ws if w[0] > rotulo[2] and w[4].upper() != "SETOR:"]
            candidato = limpar_texto(" ".join(setor_words))
            if setor_valido(candidato):
                return candidato
    return ""


def linha_tem_codigo_peca(txt: str) -> bool:
    return bool(re.fullmatch(r"[A-Z0-9]{6,}", txt)) and bool(re.search(r"\d", txt))


def linha_tem_numero_br(txt: str) -> bool:
    return bool(re.fullmatch(r"\d{1,3}(?:\.\d{3})*,\d+|\d+,\d+|\d+", txt))


def extrair_itens_por_palavras(page: fitz.Page) -> List[Tuple[str, str, float]]:
    """Extrai (codigo, descricao, qtde) usando posição X/Y das palavras do PDF."""
    words = page.get_text("words") or []
    if not words:
        return []

    # Agrupa palavras por bloco/linha quando disponível. Isso costuma respeitar as linhas da tabela.
    grupos: dict[Tuple[int, int], list] = {}
    for w in words:
        # w = (x0, y0, x1, y1, word, block_no, line_no, word_no)
        chave = (int(w[5]), int(w[6])) if len(w) >= 8 else (0, int(round(w[1])))
        grupos.setdefault(chave, []).append(w)

    itens: List[Tuple[str, str, float]] = []
    for _, ws in sorted(grupos.items(), key=lambda kv: min(w[1] for w in kv[1])):
        ws = sorted(ws, key=lambda w: w[0])
        texto_linha = " ".join(w[4] for w in ws)
        if "Total do Setor" in texto_linha or "Total Geral" in texto_linha:
            continue

        codigo = ""
        for w in ws:
            if w[0] < 95 and linha_tem_codigo_peca(w[4]):
                codigo = w[4]
                break
        if not codigo:
            continue

        qtde = None
        for w in ws:
            # No relatório, Qtde.Peças fica na primeira coluna numérica, em torno de x 320-380.
            if 280 <= w[0] <= 390 and linha_tem_numero_br(w[4]):
                qtde = numero_br_para_float(w[4])
                break
        if qtde is None:
            # Fallback: primeira quantidade depois do código.
            depois_codigo = False
            for w in ws:
                if w[4] == codigo:
                    depois_codigo = True
                    continue
                if depois_codigo and linha_tem_numero_br(w[4]):
                    qtde = numero_br_para_float(w[4])
                    break
        if qtde is None:
            continue

        desc_words = []
        for w in ws:
            palavra = w[4]
            if palavra == codigo:
                continue
            # Descrição no exemplo fica entre x 80 e x 320, antes da Qtde.Peças.
            if 70 <= w[0] < 315 and not linha_tem_numero_br(palavra):
                desc_words.append(palavra)
        descricao = limpar_texto(" ".join(desc_words))
        if not descricao:
            continue
        itens.append((codigo, descricao, qtde))

    return itens


def extrair_itens_por_linhas(page: fitz.Page) -> List[Tuple[str, str, float]]:
    """Fallback textual: usa código em uma linha, descrição na anterior e quantidade na posterior."""
    linhas = linhas_texto_pagina(page)
    itens: List[Tuple[str, str, float]] = []
    for i, linha in enumerate(linhas):
        if not linha_tem_codigo_peca(linha):
            continue
        if i == 0 or i + 1 >= len(linhas):
            continue
        descricao = linhas[i - 1]
        if descricao.upper().startswith(("PEÇA", "TOTAL", "SETOR")):
            continue
        qtde_linha = linhas[i + 1]
        if not linha_tem_numero_br(qtde_linha):
            continue
        itens.append((linha, descricao, numero_br_para_float(qtde_linha)))
    return itens


def extrair_registros_pdf(caminho_pdf: str | Path) -> List[RegistroPDF]:
    caminho_pdf = Path(caminho_pdf)
    registros: List[RegistroPDF] = []
    ultimo_lotes = extrair_lotes_nome_arquivo(caminho_pdf)
    ultimo_produto = ""
    ultimo_setor = ""

    with fitz.open(str(caminho_pdf)) as doc:
        for idx, page in enumerate(doc, start=1):
            lotes, produto_cabecalho, setor_pdf = extrair_cabecalho(page)

            # PDFs com várias páginas às vezes repetem somente a tabela nas páginas seguintes.
            # Nesses casos, reaproveitamos o cabeçalho válido da página anterior do mesmo PDF.
            if lote_valido(lotes):
                ultimo_lotes = lotes
            else:
                lotes = ultimo_lotes

            if produto_cabecalho:
                ultimo_produto = produto_cabecalho
            else:
                produto_cabecalho = ultimo_produto

            if setor_valido(setor_pdf):
                ultimo_setor = setor_pdf
            else:
                setor_pdf = ultimo_setor

            produto_equipamento = limpar_texto(setor_pdf) or "SEM_SETOR"
            operacao = produto_equipamento
            setor = produto_equipamento
            maquina_posto = produto_equipamento

            itens = extrair_itens_por_palavras(page)
            if not itens:
                itens = extrair_itens_por_linhas(page)

            for codigo, descricao, qtde in itens:
                registros.append(
                    RegistroPDF(
                        arquivo=caminho_pdf.name,
                        pagina=idx,
                        op_lote=lotes,
                        cliente_pedido=produto_cabecalho,
                        produto_equipamento=produto_equipamento,
                        codigo_peca=str(codigo).strip(),
                        descricao_peca=descricao,
                        operacao=operacao,
                        setor=setor,
                        maquina_posto=maquina_posto,
                        qtde_programada=float_para_int_se_possivel(qtde),
                    )
                )
    return registros


def celula_tem_valor(cell) -> bool:
    return cell.value not in (None, "")


def localizar_ultima_linha_dados(ws) -> int:
    """Última linha com dados nas colunas principais B, E, F ou J."""
    ultima = 1
    for row in range(ws.max_row, 1, -1):
        if any(celula_tem_valor(ws.cell(row, col)) for col in [COL_OP_LOTE, COL_CODIGO_PECA, COL_DESCRICAO_PECA, COL_QTDE_PROGRAMADA]):
            ultima = row
            break
    return ultima


def localizar_linha_modelo_formulas(ws) -> int:
    """Procura uma linha com fórmulas boas para copiar/traduzir."""
    for row in range(2, min(ws.max_row, 200) + 1):
        qtd_formulas = 0
        for col in COLUNAS_FORMULAS_PREFERIDAS:
            v = ws.cell(row, col).value
            if isinstance(v, str) and v.startswith("=") and len(v) > 1:
                qtd_formulas += 1
        if qtd_formulas >= 5:
            return row
    return max(2, localizar_ultima_linha_dados(ws))


def copiar_estilo_e_formulas(ws, linha_origem_estilo: int, linha_origem_formula: int, linha_destino: int) -> None:
    """Copia estilos da linha anterior/modelo e traduz fórmulas para a linha nova."""
    ws.row_dimensions[linha_destino].height = ws.row_dimensions[linha_origem_estilo].height

    for col in range(1, ws.max_column + 1):
        origem_estilo = ws.cell(linha_origem_estilo, col)
        destino = ws.cell(linha_destino, col)

        if origem_estilo.has_style:
            destino._style = copy.copy(origem_estilo._style)
        if origem_estilo.number_format:
            destino.number_format = origem_estilo.number_format
        if origem_estilo.alignment:
            destino.alignment = copy.copy(origem_estilo.alignment)
        if origem_estilo.protection:
            destino.protection = copy.copy(origem_estilo.protection)
        if origem_estilo.border:
            destino.border = copy.copy(origem_estilo.border)
        if origem_estilo.fill:
            destino.fill = copy.copy(origem_estilo.fill)
        if origem_estilo.font:
            destino.font = copy.copy(origem_estilo.font)

        origem_formula = ws.cell(linha_origem_formula, col)
        formula = origem_formula.value
        if isinstance(formula, str) and formula.startswith("=") and len(formula) > 1:
            try:
                destino.value = Translator(formula, origin=origem_formula.coordinate).translate_formula(destino.coordinate)
            except Exception:
                destino.value = formula


def montar_chave(registro: RegistroPDF) -> Tuple[str, str, str, str]:
    return (
        limpar_texto(registro.op_lote).upper(),
        limpar_texto(registro.produto_equipamento).upper(),
        limpar_texto(registro.codigo_peca).upper(),
        limpar_texto(registro.descricao_peca).upper(),
    )


def carregar_chaves_existentes(ws) -> set[Tuple[str, str, str, str]]:
    chaves = set()
    for row in range(2, localizar_ultima_linha_dados(ws) + 1):
        op_lote = ws.cell(row, COL_OP_LOTE).value
        produto = ws.cell(row, COL_PROD_EQUIP).value
        codigo = ws.cell(row, COL_CODIGO_PECA).value
        descricao = ws.cell(row, COL_DESCRICAO_PECA).value
        if op_lote and codigo:
            chaves.add((limpar_texto(op_lote).upper(), limpar_texto(produto).upper(), limpar_texto(codigo).upper(), limpar_texto(descricao).upper()))
    return chaves


def alimentar_planilha(
    caminho_planilha: str | Path,
    caminhos_pdfs: Iterable[str | Path],
    caminho_saida: str | Path,
    aba: str = ABA_PADRAO,
    data_programada: Optional[datetime] = None,
    evitar_duplicados: bool = True,
    preencher_cliente_pedido: bool = True,
) -> dict:
    caminho_planilha = Path(caminho_planilha)
    caminho_saida = Path(caminho_saida)
    caminhos_pdfs = [Path(p) for p in caminhos_pdfs]

    registros: List[RegistroPDF] = []
    for pdf in caminhos_pdfs:
        registros_pdf = extrair_registros_pdf(pdf)
        registros.extend(registros_pdf)

    if not registros:
        raise RuntimeError("Nenhum item foi extraído dos PDFs selecionados.")

    wb = load_workbook(caminho_planilha)
    if aba not in wb.sheetnames:
        # Caso o nome seja diferente, usa a quinta aba do arquivo.
        if len(wb.worksheets) >= 5:
            ws = wb.worksheets[4]
        else:
            raise RuntimeError(f"Aba '{aba}' não encontrada e o arquivo não tem 5 abas.")
    else:
        ws = wb[aba]

    ultima = localizar_ultima_linha_dados(ws)
    linha_modelo_formula = localizar_linha_modelo_formulas(ws)
    chaves_existentes = carregar_chaves_existentes(ws) if evitar_duplicados else set()

    inseridos = 0
    pulados = 0
    detalhes = []

    for registro in registros:
        chave = montar_chave(registro)
        if evitar_duplicados and chave in chaves_existentes:
            pulados += 1
            detalhes.append(f"PULADO duplicado: {registro.arquivo} | lote {registro.op_lote} | peça {registro.codigo_peca}")
            continue

        linha_nova = ultima + 1
        copiar_estilo_e_formulas(ws, ultima, linha_modelo_formula, linha_nova)

        if data_programada is not None:
            ws.cell(linha_nova, COL_DATA_PROGRAMADA).value = data_programada

        ws.cell(linha_nova, COL_OP_LOTE).value = registro.op_lote
        if preencher_cliente_pedido:
            ws.cell(linha_nova, COL_CLIENTE_PEDIDO).value = registro.cliente_pedido
        ws.cell(linha_nova, COL_PROD_EQUIP).value = registro.produto_equipamento
        ws.cell(linha_nova, COL_CODIGO_PECA).value = registro.codigo_peca
        ws.cell(linha_nova, COL_DESCRICAO_PECA).value = registro.descricao_peca
        ws.cell(linha_nova, COL_OPERACAO).value = registro.operacao
        ws.cell(linha_nova, COL_SETOR).value = registro.setor
        ws.cell(linha_nova, COL_MAQUINA_POSTO).value = registro.maquina_posto
        ws.cell(linha_nova, COL_QTDE_PROGRAMADA).value = registro.qtde_programada

        # Fórmulas auxiliares úteis para a aba.
        # Se a fórmula já foi copiada do modelo, estas linhas preservam a lógica comum.
        # Ajuste/remova se sua planilha tiver outra lógica nessas colunas.
        if ws.cell(linha_nova, 28).value in (None, ""):
            ws.cell(linha_nova, 28).value = f'=B{linha_nova}&"|"&I{linha_nova}'  # AB
        if ws.cell(linha_nova, 29).value in (None, ""):
            ws.cell(linha_nova, 29).value = f'=IF($B{linha_nova}="","","|"&SUBSTITUTE(TRIM($B{linha_nova}&"")," ","|")&"|")'  # AC
        if ws.cell(linha_nova, 30).value in (None, ""):
            ws.cell(linha_nova, 30).value = f'=TRIM($I{linha_nova}&"")'  # AD

        chaves_existentes.add(chave)
        ultima = linha_nova
        inseridos += 1
        detalhes.append(f"INSERIDO: {registro.arquivo} | lote {registro.op_lote} | peça {registro.codigo_peca} | qtd {registro.qtde_programada}")

    # Força recálculo no Excel ao abrir.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)
    return {
        "saida": str(caminho_saida),
        "pdfs": len(caminhos_pdfs),
        "registros_extraidos": len(registros),
        "inseridos": inseridos,
        "pulados_duplicados": pulados,
        "detalhes": detalhes,
    }


def gerar_nome_saida(caminho_planilha: str | Path) -> str:
    p = Path(caminho_planilha)
    carimbo = datetime.now().strftime("%Y%m%d_%H%M")
    return str(p.with_name(f"{p.stem}_atualizada_{carimbo}{p.suffix}"))


def rodar_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox
    from tkinter import ttk

    root = tk.Tk()
    root.title("Alimentador PPCP por PDFs")
    root.geometry("920x620")

    planilha_var = tk.StringVar()
    saida_var = tk.StringVar()
    aba_var = tk.StringVar(value=ABA_PADRAO)
    data_var = tk.StringVar()
    duplicados_var = tk.BooleanVar(value=True)
    cliente_var = tk.BooleanVar(value=True)
    pdfs: List[str] = []

    def escolher_planilha():
        arq = filedialog.askopenfilename(
            title="Selecione a planilha mãe",
            filetypes=[("Excel", "*.xlsx"), ("Todos os arquivos", "*.*")],
        )
        if arq:
            planilha_var.set(arq)
            saida_var.set(gerar_nome_saida(arq))

    def escolher_pdfs():
        selecionados = filedialog.askopenfilenames(
            title="Selecione os PDFs",
            filetypes=[("PDF", "*.pdf"), ("Todos os arquivos", "*.*")],
        )
        if selecionados:
            pdfs.clear()
            pdfs.extend(selecionados)
            lista_pdfs.delete(0, tk.END)
            for p in pdfs:
                lista_pdfs.insert(tk.END, p)

    def escolher_saida():
        arq = filedialog.asksaveasfilename(
            title="Salvar planilha atualizada como",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=Path(saida_var.get()).name if saida_var.get() else "PPCP_ACOMPAN_atualizada.xlsx",
        )
        if arq:
            saida_var.set(arq)

    def log(msg: str):
        texto_log.insert(tk.END, msg + "\n")
        texto_log.see(tk.END)
        root.update_idletasks()

    def processar():
        try:
            if not planilha_var.get():
                messagebox.showwarning("Atenção", "Selecione a planilha mãe.")
                return
            if not pdfs:
                messagebox.showwarning("Atenção", "Selecione um ou mais PDFs.")
                return
            if not saida_var.get():
                saida_var.set(gerar_nome_saida(planilha_var.get()))

            data_prog = parse_data_opcional(data_var.get())
            texto_log.delete("1.0", tk.END)
            log("Processando PDFs...")

            resultado = alimentar_planilha(
                caminho_planilha=planilha_var.get(),
                caminhos_pdfs=pdfs,
                caminho_saida=saida_var.get(),
                aba=aba_var.get() or ABA_PADRAO,
                data_programada=data_prog,
                evitar_duplicados=duplicados_var.get(),
                preencher_cliente_pedido=cliente_var.get(),
            )
            for d in resultado["detalhes"]:
                log(d)
            log("-" * 70)
            log(f"PDFs lidos: {resultado['pdfs']}")
            log(f"Registros extraídos: {resultado['registros_extraidos']}")
            log(f"Linhas inseridas: {resultado['inseridos']}")
            log(f"Duplicados pulados: {resultado['pulados_duplicados']}")
            log(f"Arquivo salvo em: {resultado['saida']}")
            messagebox.showinfo("Concluído", f"Planilha gerada com sucesso:\n{resultado['saida']}")
        except Exception as e:
            messagebox.showerror("Erro", str(e))
            log(f"ERRO: {e}")

    frame = ttk.Frame(root, padding=12)
    frame.pack(fill=tk.BOTH, expand=True)

    ttk.Label(frame, text="Planilha mãe:").grid(row=0, column=0, sticky="w")
    ttk.Entry(frame, textvariable=planilha_var, width=95).grid(row=0, column=1, sticky="we", padx=6)
    ttk.Button(frame, text="Selecionar", command=escolher_planilha).grid(row=0, column=2, sticky="e")

    ttk.Label(frame, text="PDFs:").grid(row=1, column=0, sticky="nw", pady=(8, 0))
    lista_pdfs = tk.Listbox(frame, height=6)
    lista_pdfs.grid(row=1, column=1, sticky="nsew", padx=6, pady=(8, 0))
    ttk.Button(frame, text="Selecionar PDFs", command=escolher_pdfs).grid(row=1, column=2, sticky="ne", pady=(8, 0))

    ttk.Label(frame, text="Aba destino:").grid(row=2, column=0, sticky="w", pady=(8, 0))
    ttk.Entry(frame, textvariable=aba_var, width=30).grid(row=2, column=1, sticky="w", padx=6, pady=(8, 0))

    ttk.Label(frame, text="Data Programada opcional:").grid(row=3, column=0, sticky="w", pady=(8, 0))
    ttk.Entry(frame, textvariable=data_var, width=30).grid(row=3, column=1, sticky="w", padx=6, pady=(8, 0))
    ttk.Label(frame, text="Ex.: 30/06/2026. Se deixar vazio, a coluna A fica vazia nas linhas novas.").grid(row=3, column=1, sticky="w", padx=220, pady=(8, 0))

    ttk.Label(frame, text="Salvar como:").grid(row=4, column=0, sticky="w", pady=(8, 0))
    ttk.Entry(frame, textvariable=saida_var, width=95).grid(row=4, column=1, sticky="we", padx=6, pady=(8, 0))
    ttk.Button(frame, text="Alterar", command=escolher_saida).grid(row=4, column=2, sticky="e", pady=(8, 0))

    ttk.Checkbutton(frame, text="Evitar duplicados", variable=duplicados_var).grid(row=5, column=1, sticky="w", padx=6, pady=(8, 0))
    ttk.Checkbutton(frame, text="Preencher Cliente/Pedido com produto do cabeçalho do PDF", variable=cliente_var).grid(row=6, column=1, sticky="w", padx=6, pady=(4, 0))

    ttk.Button(frame, text="PROCESSAR", command=processar).grid(row=7, column=1, sticky="w", padx=6, pady=12)

    ttk.Label(frame, text="Log:").grid(row=8, column=0, sticky="nw")
    texto_log = tk.Text(frame, height=14)
    texto_log.grid(row=8, column=1, columnspan=2, sticky="nsew", padx=6)

    frame.columnconfigure(1, weight=1)
    frame.rowconfigure(8, weight=1)
    root.mainloop()


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Alimenta a aba 5_ACOMPANHAMENTO com dados extraídos de PDFs.")
    parser.add_argument("--sem-gui", action="store_true", help="Executa sem abrir interface.")
    parser.add_argument("--planilha", help="Caminho da planilha mãe .xlsx")
    parser.add_argument("--pdfs", nargs="*", help="Lista de PDFs")
    parser.add_argument("--saida", help="Caminho do arquivo .xlsx de saída")
    parser.add_argument("--aba", default=ABA_PADRAO, help="Nome da aba destino")
    parser.add_argument("--data", default="", help="Data Programada DD/MM/AAAA opcional")
    parser.add_argument("--permitir-duplicados", action="store_true", help="Não pula registros duplicados")
    parser.add_argument("--nao-preencher-cliente", action="store_true", help="Não preenche a coluna Cliente/Pedido")
    args = parser.parse_args(argv)

    if not args.sem_gui:
        rodar_gui()
        return 0

    if not args.planilha or not args.pdfs:
        parser.error("No modo --sem-gui informe --planilha e --pdfs")

    saida = args.saida or gerar_nome_saida(args.planilha)
    data_prog = parse_data_opcional(args.data)
    resultado = alimentar_planilha(
        caminho_planilha=args.planilha,
        caminhos_pdfs=args.pdfs,
        caminho_saida=saida,
        aba=args.aba,
        data_programada=data_prog,
        evitar_duplicados=not args.permitir_duplicados,
        preencher_cliente_pedido=not args.nao_preencher_cliente,
    )
    print("Concluído.")
    print(f"PDFs lidos: {resultado['pdfs']}")
    print(f"Registros extraídos: {resultado['registros_extraidos']}")
    print(f"Linhas inseridas: {resultado['inseridos']}")
    print(f"Duplicados pulados: {resultado['pulados_duplicados']}")
    print(f"Arquivo salvo em: {resultado['saida']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())